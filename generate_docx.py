import sys
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_ORIENT

def create_doc():
    doc = Document()
    
    # Set orientation to Landscape for wider table
    section = doc.sections[-1]
    new_width, new_height = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_width
    section.page_height = new_height
    
    # Title
    doc.add_heading('Dokumen Blackbox Testing - Register Page', 0)
    
    p = doc.add_paragraph()
    p.add_run('Aplikasi: ').bold = True
    p.add_run('Beaply — Scholarship Insight\n')
    p.add_run('Modul: ').bold = True
    p.add_run('Register Page (Sign Up)\n')
    p.add_run('Tanggal Uji: ').bold = True
    p.add_run('25 Mei 2026\n')
    p.add_run('Status Keseluruhan: ').bold = True
    p.add_run('17 Test Cases PASS\n')
    
    # Main Table
    table = doc.add_table(rows=1, cols=9)
    table.style = 'Table Grid'
    
    headers = [
        "ID Test", 
        "Skenario Pengujian", 
        "Langkah-langkah", 
        "Data Uji", 
        "Ekspektasi Hasil", 
        "Gambar", 
        "Hasil Aktual", 
        "Status", 
        "Komentar"
    ]
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        
    data = [
        ["TC-AUTH-01", "Menguji proses registrasi dengan semua data wajib valid dan benar.", "1. Buka aplikasi.\n2. Ke tab Sign up.\n3. Isi semua field dengan data valid.\n4. Klik Save Profile.", "Nama: Tutut\nEmail: user@mail.com\nPass: Test123!\nDOB: 2005-12-30", "Sistem menyimpan akun & profil dan redirect ke Dashboard.", "", "Registrasi berhasil. Akun tersimpan. Redirect sukses.", "Pass", "Lolos validasi auth dan profil"],
        ["TC-AUTH-02", "Menguji proses registrasi dengan email yang sudah terdaftar.", "1. Buka aplikasi.\n2. Ke tab Sign up.\n3. Masukkan email yang sudah ada.\n4. Klik Save Profile.", "Email: user@mail.com\nPass: Test456@", "Sistem menolak dan menampilkan pesan Email sudah terdaftar.", "", "Sistem menolak dengan pesan error yang sesuai.", "Pass", "Pengecekan duplikasi email berhasil"],
        ["TC-AUTH-03", "Data tidak diisi apapun", "1. Buka aplikasi.\n2. Ke tab Sign up.\n3. Biarkan field kosong.\n4. Klik Save Profile.", "Kosong", "Sistem menolak dan meminta field wajib diisi.", "", "Sistem menolak dan menampilkan error validasi.", "Pass", ""],
        ["TC-AUTH-04", "Data (*) diisi salah satu saja", "1. Buka aplikasi.\n2. Isi field Nama saja.\n3. Klik Save Profile.", "Nama: Tutut\nLainnya kosong", "Sistem menolak dan meminta field wajib lain diisi.", "", "Sistem menolak. Pesan format email muncul karena kosong.", "Pass", ""],
        ["TC-AUTH-05", "Email diisi tanpa @", "1. Buka aplikasi.\n2. Masukkan email tanpa @.\n3. Klik Save Profile.", "Email: tutgmail.com", "Sistem menolak karena format email tidak valid.", "", "Sistem menolak. Regex memvalidasi tidak ada @.", "Pass", ""],
        ["TC-AUTH-06", "Email diisi tanpa TLD (.com)", "1. Buka aplikasi.\n2. Masukkan email tanpa TLD.\n3. Klik Save Profile.", "Email: tut@gmail", "Sistem menolak format email.", "", "Sistem menolak karena regex butuh ekstensi TLD.", "Pass", ""],
        ["TC-AUTH-07", "Password tanpa simbol khusus", "1. Buka aplikasi.\n2. Masukkan password tanpa simbol.\n3. Klik Save Profile.", "Pass: Tut123456789", "Sistem menolak, password harus mengandung simbol.", "", "Sistem menolak, pesan error syarat simbol muncul.", "Pass", ""],
        ["TC-AUTH-08", "Password tanpa huruf besar", "1. Buka aplikasi.\n2. Masukkan password tanpa huruf kapital.\n3. Klik Save Profile.", "Pass: tutut123!", "Sistem menolak, password harus ada huruf besar.", "", "Sistem menolak, pesan error huruf besar muncul.", "Pass", ""],
        ["TC-AUTH-09", "Password tanpa angka", "1. Buka aplikasi.\n2. Masukkan password tanpa angka.\n3. Klik Save Profile.", "Pass: Tutut!!!", "Sistem menolak, password harus mengandung angka.", "", "Sistem menolak, pesan error angka muncul.", "Pass", ""],
        ["TC-AUTH-10", "Hari pada tanggal lahir > 31", "1. Buka aplikasi.\n2. Masukkan tanggal dengan hari 35.\n3. Klik Save Profile.", "DOB: 2005-12-35", "Sistem menolak dengan error format tanggal salah.", "", "Sistem menolak, input tidak valid secara kalender.", "Pass", "Hari > 31 otomatis ditolak oleh sistem"],
        ["TC-AUTH-11", "Nama kurang dari 2 karakter", "1. Buka aplikasi.\n2. Masukkan nama 1 huruf.\n3. Klik Save Profile.", "Nama: T", "Sistem menolak dan meminta minimal 2 karakter.", "", "Sistem menolak dengan pesan yang sesuai.", "Pass", ""],
        ["TC-AUTH-12", "GPA di luar range (0.00 - 4.00)", "1. Buka aplikasi.\n2. Masukkan GPA melebihi batas.\n3. Klik Save Profile.", "GPA: 4.50", "Sistem menolak GPA > 4.00.", "", "Sistem menolak karena nilai 4.50 di luar rentang.", "Pass", ""],
        ["TC-AUTH-13", "Semester = 0", "1. Buka aplikasi.\n2. Masukkan semester = 0.\n3. Klik Save Profile.", "Semester: 0", "Sistem menolak semester 0.", "", "Sistem menolak nilai 0, harus minimal 1.", "Pass", ""],
        ["TC-AUTH-14", "IELTS > 9.0", "1. Buka aplikasi.\n2. Masukkan IELTS melebihi batas.\n3. Klik Save Profile.", "IELTS: 10.0", "Sistem menolak nilai IELTS.", "", "Sistem menolak nilai 10.0 (melebihi 9.0).", "Pass", "Data opsional tervalidasi jika diisi"],
        ["TC-AUTH-15", "TOEFL > 120", "1. Buka aplikasi.\n2. Masukkan TOEFL melebihi batas.\n3. Klik Save Profile.", "TOEFL: 130", "Sistem menolak nilai TOEFL.", "", "Sistem menolak nilai 130 (melebihi 120).", "Pass", ""],
        ["TC-AUTH-16", "Format tanggal salah", "1. Buka aplikasi.\n2. Masukkan format DD-MM-YYYY.\n3. Klik Save Profile.", "DOB: 30-12-2005", "Sistem menolak format tanggal salah.", "", "Sistem menolak input DD-MM-YYYY.", "Pass", "Mengharuskan format standar (YYYY-MM-DD)"],
        ["TC-AUTH-17", "Data valid lengkap + data opsional", "1. Buka aplikasi.\n2. Isi semua data valid dan skor opsional.\n3. Klik Save Profile.", "KIP: Ya\nIELTS: 7.5\nTOEFL: 100\nDuolingo: 120", "Sistem menyimpan semua data (wajib & opsional).", "", "Registrasi sukses, data opsional ikut tersimpan dengan benar.", "Pass", "Tes memastikan data opsional valid diproses lancar"]
    ]
    
    for row_data in data:
        row_cells = table.add_row().cells
        for i, text in enumerate(row_data):
            row_cells[i].text = text
            
    # Set styling font untuk sel tabel
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    
    # Set bold untuk header
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                
    doc.save("Dokumen_Blackbox_Testing_Register.docx")
    print("Dokumen berhasil dibuat: Dokumen_Blackbox_Testing_Register.docx")

if __name__ == "__main__":
    create_doc()
